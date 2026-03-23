"""Markdown -> DOCX 导出工具。

当前实现优先保证企业内常见的标题、列表和代码块样式可读，避免为了复杂富文本把导出链路做得过重。
"""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.shared import Pt


def markdown_to_docx_bytes(title: str, markdown_text: str) -> bytes:
    """将 PRD Markdown 转成可下载的 Word 二进制内容。"""
    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Microsoft YaHei"
    normal_style.font.size = Pt(10.5)

    document.add_heading(title, level=0)

    in_code_block = False
    for raw_line in markdown_text.splitlines():
        line = raw_line.rstrip()
        if line.startswith("```"):
            in_code_block = not in_code_block
            continue
        if not line:
            document.add_paragraph("")
            continue
        if in_code_block:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            continue
        if line.startswith("# "):
            document.add_heading(line[2:].strip(), level=1)
            continue
        if line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
            continue
        if line.startswith("### "):
            document.add_heading(line[4:].strip(), level=3)
            continue
        if line.startswith("- "):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
            continue
        if line[:2].isdigit() and line[1] == ".":
            document.add_paragraph(line[2:].strip(), style="List Number")
            continue
        document.add_paragraph(line)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
