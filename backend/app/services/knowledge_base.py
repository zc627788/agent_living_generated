"""知识库管理服务。

负责文档入库、版本快照、切片、向量写入和内置知识种子初始化，是 RAG 的数据入口。
"""

from __future__ import annotations

import mimetypes
import re
import uuid
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from qdrant_client.models import PointStruct

from app.core.config import settings
from app.db.database import database
from app.models.schemas import (
    KnowledgeChunkPreview,
    KnowledgeDocumentDetail,
    KnowledgeDocumentSummary,
    KnowledgeDocumentUpdateRequest,
    KnowledgeDocumentVersionDetail,
    KnowledgeDocumentVersionSummary,
)
from app.services.llm_service import llm_service
from app.services.object_storage import object_storage_service
from app.services.vector_store import vector_store_service

CHUNK_SIZE = 1000
CHUNK_OVERLAP = 200
SUPPORTED_SUFFIXES = {".md", ".txt", ".docx"}


@dataclass
class ParsedDocument:
    title: str
    source_type: str
    source: str
    url: str | None
    thumbnail_url: str | None
    tags: list[str]
    body: str


class KnowledgeBaseService:
    def __init__(self, built_in_dir: Path) -> None:
        self.built_in_dir = built_in_dir

    def refresh(self) -> tuple[int, int]:
        database.initialize()
        self._seed_builtin_documents()
        records = database.list_document_records()
        return self._reindex_records(records, recreate_collection=True)

    def ensure_seeded(self) -> tuple[int, int]:
        database.initialize()
        self._seed_builtin_documents()
        documents, chunks = database.knowledge_counts()
        if documents == 0 or chunks == 0 or not vector_store_service.collection_exists():
            return self.refresh()
        return documents, chunks

    def list_documents(self) -> list[KnowledgeDocumentSummary]:
        database.initialize()
        return database.list_documents()

    def get_document(self, document_id: str) -> KnowledgeDocumentDetail | None:
        database.initialize()
        return database.get_document(document_id)

    def list_chunks(self, document_id: str | None = None) -> list[KnowledgeChunkPreview]:
        database.initialize()
        return database.list_chunks(document_id)

    def save_upload(
        self,
        *,
        file_name: str,
        content: bytes,
        title: str,
        source_type: str,
        source: str,
        tags: list[str],
        url: str | None,
        thumbnail_url: str | None,
    ) -> KnowledgeDocumentDetail:
        """保存上传文档、写入对象存储，并立即触发该文档的索引重建。"""
        suffix = Path(file_name).suffix.lower()
        if suffix not in SUPPORTED_SUFFIXES:
            raise ValueError("仅支持 .md / .txt / .docx 文件。")

        body = self._normalize_upload_content(content, suffix)
        object_key = f"knowledge/uploads/{uuid.uuid4().hex}/{Path(file_name).name}"
        object_storage_service.upload_bytes(object_key, content, self._content_type(file_name))

        document_id = database.create_document(
            title=title,
            source_type=source_type,
            source=source,
            url=url,
            thumbnail_url=thumbnail_url,
            tags=tags,
            content=body,
            storage_key=object_key,
            origin_file_name=file_name,
            is_builtin=False,
        )
        database.create_document_version_snapshot(document_id, "初始上传")
        self.reindex_documents([document_id])
        return self._require_document(document_id)

    def update_document(self, document_id: str, payload: KnowledgeDocumentUpdateRequest) -> KnowledgeDocumentDetail:
        document = self._require_document(document_id)
        if document.is_builtin:
            raise ValueError("内置文档不支持在线编辑，请修改内置知识源文件后再重建索引。")

        database.update_document(
            document_id,
            title=payload.title,
            source_type=payload.source_type,
            source=payload.source,
            url=payload.url,
            thumbnail_url=payload.thumbnail_url,
            tags=payload.tags,
            content=payload.content,
        )

        if document.storage_key:
            object_storage_service.upload_bytes(
                document.storage_key,
                payload.content.encode("utf-8"),
                "text/markdown; charset=utf-8",
            )

        database.create_document_version_snapshot(document_id, "手动编辑")
        self.reindex_documents([document_id])
        return self._require_document(document_id)

    def delete_document(self, document_id: str) -> tuple[int, int]:
        document = self._require_document(document_id)
        if document.is_builtin:
            raise ValueError("内置文档不支持删除。")

        vector_store_service.delete_points_by_document_ids([document_id])
        if document.storage_key:
            object_storage_service.delete_object(document.storage_key)
        database.delete_document(document_id)
        return database.knowledge_counts()

    def reindex_documents(self, document_ids: list[str]) -> tuple[int, int]:
        database.initialize()
        if not document_ids:
            return database.knowledge_counts()

        for document_id in document_ids:
            self._require_document(document_id)

        self._seed_builtin_documents()
        if not vector_store_service.collection_exists():
            return self.refresh()

        records = database.list_document_records(document_ids)
        return self._reindex_records(records, recreate_collection=False)

    def list_document_versions(self, document_id: str) -> list[KnowledgeDocumentVersionSummary]:
        self._require_document(document_id)
        return database.list_document_versions(document_id)

    def restore_document_version(self, document_id: str, version_id: str) -> KnowledgeDocumentDetail:
        document = self._require_document(document_id)
        if document.is_builtin:
            raise ValueError("内置文档不支持版本恢复。")

        version = database.get_document_version(version_id)
        if version is None or version.document_id != document_id:
            raise ValueError("版本不存在。")

        database.update_document(
            document_id,
            title=version.title,
            source_type=version.source_type,
            source=version.source,
            url=version.url,
            thumbnail_url=version.thumbnail_url,
            tags=version.tags,
            content=version.content,
            storage_key=version.storage_key,
            origin_file_name=version.origin_file_name,
        )

        if version.storage_key:
            object_storage_service.upload_bytes(
                version.storage_key,
                version.content.encode("utf-8"),
                "text/markdown; charset=utf-8",
            )

        database.create_document_version_snapshot(document_id, f"恢复到 V{version.version_number}")
        self.reindex_documents([document_id])
        return self._require_document(document_id)

    def _seed_builtin_documents(self) -> None:
        if not self.built_in_dir.exists():
            return

        for path in sorted(self.built_in_dir.glob("*.md")):
            parsed = self._parse_markdown_document(path)
            database.upsert_builtin_document(
                title=parsed.title,
                source_type=parsed.source_type,
                source=parsed.source,
                url=parsed.url,
                thumbnail_url=parsed.thumbnail_url,
                tags=parsed.tags,
                content=parsed.body,
                storage_key=f"builtin/{path.name}",
            )

    def _reindex_records(self, records: list[dict], *, recreate_collection: bool) -> tuple[int, int]:
        if recreate_collection:
            database.delete_chunks_for_documents()
        else:
            document_ids = [str(record["id"]) for record in records]
            database.delete_chunks_for_documents(document_ids)
            vector_store_service.delete_points_by_document_ids(document_ids)

        embedding_inputs: list[str] = []
        point_payloads: list[dict] = []

        for record in records:
            document_id = str(record["id"])
            tags = self._coerce_tags(record.get("tags_json", []))
            chunks = self._chunk_document(str(record["content"]))
            chunk_ids = [str(uuid.uuid4()) for _ in chunks]
            for index, chunk in enumerate(chunks):
                prev_chunk_id = None if index == 0 else chunk_ids[index - 1]
                next_chunk_id = None if index == len(chunks) - 1 else chunk_ids[index + 1]
                database.insert_chunk(
                    chunk_id=chunk_ids[index],
                    document_id=document_id,
                    chunk_index=index,
                    section_path=chunk["section_path"],
                    content=chunk["content"],
                    prev_chunk_id=prev_chunk_id,
                    next_chunk_id=next_chunk_id,
                    tags=tags,
                )
                embedding_inputs.append(chunk["content"])
                point_payloads.append(
                    {
                        "id": chunk_ids[index],
                        "document_id": document_id,
                        "title": str(record["title"]),
                        "source": str(record["source"]),
                        "source_type": str(record["source_type"]),
                        "url": record["url"],
                        "thumbnail_url": record["thumbnail_url"],
                        "section_path": chunk["section_path"],
                        "tags": tags,
                        "content": chunk["content"],
                        "chunk_index": index,
                    }
                )

        if not embedding_inputs:
            if recreate_collection:
                vector_store_service.delete_collection()
            return database.knowledge_counts()

        vectors = llm_service.embed_texts(embedding_inputs)
        if recreate_collection:
            vector_store_service.recreate_collection(len(vectors[0]))
        else:
            vector_store_service.ensure_collection(len(vectors[0]))

        points = [
            PointStruct(id=payload["id"], vector=vector, payload=payload)
            for payload, vector in zip(point_payloads, vectors, strict=False)
        ]
        vector_store_service.upsert_points(points)
        return database.knowledge_counts()

    def _require_document(self, document_id: str) -> KnowledgeDocumentDetail:
        document = database.get_document(document_id)
        if document is None:
            raise ValueError("文档不存在。")
        return document

    def _parse_markdown_document(self, path: Path) -> ParsedDocument:
        raw = self._read_text(path)
        meta_lines: list[str] = []
        body_lines: list[str] = []
        reading_meta = True

        for line in raw.splitlines():
            if reading_meta and ":" in line and line.strip():
                meta_lines.append(line.strip())
                continue
            if reading_meta and not line.strip():
                reading_meta = False
                continue
            reading_meta = False
            body_lines.append(line)

        meta: dict[str, str] = {}
        for line in meta_lines:
            key, value = line.split(":", 1)
            meta[key.strip().lower()] = value.strip()

        return ParsedDocument(
            title=meta.get("title", path.stem),
            source_type=meta.get("sourcetype", "knowledge"),
            source=meta.get("source", path.stem),
            url=meta.get("url") or None,
            thumbnail_url=meta.get("thumbnailurl") or None,
            tags=[tag.strip() for tag in meta.get("tags", "").split(",") if tag.strip()],
            body="\n".join(body_lines).strip(),
        )

    def _normalize_upload_content(self, content: bytes, suffix: str) -> str:
        if suffix == ".docx":
            document = DocxDocument(BytesIO(content))
            paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            return "\n".join(paragraphs)
        return self._decode_text_bytes(content)

    def _content_type(self, file_name: str) -> str:
        suffix = Path(file_name).suffix.lower()
        if suffix == ".md":
            return "text/markdown; charset=utf-8"
        guessed, _ = mimetypes.guess_type(file_name)
        return guessed or "application/octet-stream"

    def _read_text(self, path: Path) -> str:
        return self._decode_text_bytes(path.read_bytes())

    def _decode_text_bytes(self, content: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise ValueError("文档编码无法识别，请使用 UTF-8、UTF-8 BOM 或 GB18030 文本。")

    def _coerce_tags(self, raw_tags: object) -> list[str]:
        if isinstance(raw_tags, list):
            return [str(tag) for tag in raw_tags]
        if isinstance(raw_tags, str):
            text = raw_tags.strip()
            if text.startswith("["):
                try:
                    import json

                    data = json.loads(text)
                    if isinstance(data, list):
                        return [str(tag) for tag in data]
                except ValueError:
                    return [tag.strip() for tag in text.split(",") if tag.strip()]
            return [tag.strip() for tag in text.split(",") if tag.strip()]
        return []

    def _chunk_document(self, body: str) -> list[dict[str, str]]:
        sections = self._split_sections(body)
        chunks: list[dict[str, str]] = []

        for section_path, section_content in sections:
            text = section_content.strip()
            if not text:
                continue

            start = 0
            while start < len(text):
                end = min(start + CHUNK_SIZE, len(text))
                chunk_text = text[start:end].strip()
                if chunk_text:
                    chunks.append({"section_path": section_path, "content": chunk_text})
                if end >= len(text):
                    break
                start = max(0, end - CHUNK_OVERLAP)

        return chunks

    def _split_sections(self, body: str) -> list[tuple[str, str]]:
        sections: list[tuple[str, str]] = []
        current_path = "概览"
        buffer: list[str] = []
        heading_stack: list[str] = []

        for line in body.splitlines():
            heading_match = re.match(r"^(#{1,3})\s+(.*)$", line.strip())
            if heading_match:
                if buffer:
                    sections.append((current_path, "\n".join(buffer).strip()))
                    buffer = []
                level = len(heading_match.group(1))
                title = heading_match.group(2).strip()
                heading_stack = heading_stack[: level - 1]
                heading_stack.append(title)
                current_path = " / ".join(heading_stack)
                continue
            buffer.append(line)

        if buffer:
            sections.append((current_path, "\n".join(buffer).strip()))

        return sections


knowledge_base_service = KnowledgeBaseService(settings.knowledge_dir)
